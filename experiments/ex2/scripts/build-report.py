from pathlib import Path
from copy import deepcopy
import hashlib, json, zipfile
from PIL import Image
from docx import Document
from docx.table import Table
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(__file__).resolve().parents[3]; E=R/'experiments/ex2'; T=R/'tmp/pdfs/ex2/report'; T.mkdir(parents=True,exist_ok=True)
ref=R/'tmp/pdfs/ex1-template/实验报告模板.docx'
doc=Document(ref); template=deepcopy(doc.tables[0]._tbl); body=doc._element.body
hsh=hashlib.sha256((R/'experiments/实验报告模板.doc').read_bytes()).hexdigest()
(T/'artifact.md').write_text(f'''# Template contract
Source: experiments/实验报告模板.doc; SHA256 {hsh}
Retained converted reference: {ref}
Preserve masthead, black table form, page geometry, label width, metadata and teacher evaluation.
Clone continuation row 4 for evidence pages. Relax empty minimum heights; use exact 15pt body line spacing and no grid snapping. Pictures use automatic line spacing.
Fill known identity 易治行 2023141220023 计算机科学与技术 102班. Date 2026-10-20. Week 7 follows the student's edited experiment 1 report. Unknown hours and teacher evaluation remain blank.
Evidence: actual Packet Tracer 9.0.1 and formal dumpcap en0 capture, final/ path. No invented success for real scu ping. Figures are genuine screenshots, cropped without changing contents.
Output: DOCX and PDF. Inspect every rendered page before delivery.
''')
for el in list(body):
    if el.tag!=qn('w:sectPr') and el not in [doc.paragraphs[0]._p,doc.paragraphs[1]._p,doc.paragraphs[2]._p,doc.tables[0]._tbl]:body.remove(el)
t=doc.tables[0]
for tr in list(t._tbl.tr_lst)[4:]:t._tbl.remove(tr)

def fmt(p,bold=False,size=10.5,align=WD_ALIGN_PARAGRAPH.LEFT):
    pp=p._p.get_or_add_pPr();snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'false');pp.append(snap)
    pf=p.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(5);pf.line_spacing=Pt(15)
    pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.keep_with_next=False;pf.keep_together=False
    p.alignment=align
    for r in p.runs:
        r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Songti SC')
    return p
def clear(c):
    for el in list(c._tc):
        if el.tag!=qn('w:tcPr'):c._tc.remove(el)
    c.add_paragraph();c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    mar=OxmlElement('w:tcMar')
    for tag,v in [('top',90),('bottom',80),('left',100),('right',100)]:
        el=OxmlElement('w:'+tag);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');mar.append(el)
    c._tc.get_or_add_tcPr().append(mar)
def para(c,text,bold=False,size=10.5):
    p=c.paragraphs[0] if len(c.paragraphs)==1 and not c.paragraphs[0].text and len(c.paragraphs[0]._p)==0 else c.add_paragraph()
    p.add_run(text);return fmt(p,bold,size)
def fill(c,lines):
    clear(c)
    for s in lines:para(c,s)
def clean(row):
    pr=row._tr.get_or_add_trPr()
    for el in list(pr):
        if el.tag in [qn('w:trHeight'),qn('w:cantSplit')]:pr.remove(el)
def page(label,title):
    p=doc.add_paragraph();p.paragraph_format.page_break_before=True;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
    el=deepcopy(template)
    for tr in list(el.tr_lst):el.remove(tr)
    el.append(deepcopy(template.tr_lst[4]));body.insert(len(body)-1,el)
    tt=Table(el,doc._body);row=tt.rows[0];clean(row)
    for c in [row.cells[0],row.cells[1]]:clear(c)
    para(row.cells[0],label);row.cells[0].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in row.cells[0].paragraphs:fmt(p,align=WD_ALIGN_PARAGRAPH.CENTER)
    c=row.cells[1];para(c,title,True,11);return tt,c
fig=0
def pic(c,name,caption,crop=None,width=4.65):
    global fig
    src=E/'evidence'/name; im=Image.open(src).convert('RGB')
    if crop:im=im.crop(crop)
    out=T/(src.stem+'.png');im.save(out)
    p=c.add_paragraph();fmt(p,align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(out),width=Inches(width));fig+=1
    p=para(c,f'图{fig} {caption}',size=9);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def qa(c,n,q,a):
    p=para(c,f'{n}. {q}',True);p.paragraph_format.keep_with_next=True;para(c,a)
def mono(c,s):
    p=para(c,s,size=9)
    for r in p.runs:r.font.name='Menlo'

p=doc.paragraphs[2];p.text='学号：2023141220023  姓名：易治行\n专业：计算机科学与技术  班级：102班  第7周';fmt(p)
fill(t.cell(0,1),['计算机网络']);fill(t.cell(0,3),[''])
fill(t.cell(1,1),['实验2 DNS服务器配置与DNS协议分析']);fill(t.cell(1,3),['2026年10月20日'])
fill(t.cell(2,1),['1. 在 Packet Tracer 中配置 DNS 的 A 与 CNAME 记录，验证域名解析与 IP 通信之间的关系。','2. 通过 DNS 服务器改址故障及恢复实验，掌握客户端解析器配置的排查方法。','3. 使用 Wireshark 捕获真实 DNS 查询与响应，分析事务 ID、端口、递归标志、资源记录和缓存行为。'])
fill(t.cell(3,1),['硬件与系统：Apple M2 Mac，macOS 26.5.1。','仿真：Cisco Packet Tracer 9.0.1.0858；PC-PT×1、Server-PT×2、Switch-PT×1，三条直通铜缆。','抓包：Wireshark dumpcap 4.6.9，en0 接口；TShark 4.6.9 解析；Wireshark 4.4.14 查看截图。','真实客户端：10.135.128.230。ping 抓包中实际 DNS 端点为223.5.5.5；四组 nslookup 指定同一解析器202.115.39.9。'])
for row in t.rows:
    clean(row)
    for c in row.cells:
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

_,c=page('实验内容\n（步骤和方法）','一 Packet Tracer 网络与记录配置')
para(c,'搭建同一交换机连接的局域网。PC 为192.168.1.155/24，DNS服务器初始为192.168.1.199/24，目标主机为192.168.1.200/24；三者均无需经网关进行本子网通信。PC 初始 DNS 指向192.168.1.199。')
pic(c,'pt-topology.jpg','Packet Tracer 实际拓扑与三条直通链路',(440,480,1700,1240))
para(c,'设备按学号命名：PC-2023141220023、DNS-2023141220023、HOST-2023141220023、SW-2023141220023。')
para(c,'DNS 服务开启，配置两条记录：',True)
mono(c,'www.yizhixing.com  A  192.168.1.200\nwww1.2023141220023.com  CNAME  www.yizhixing.com')
para(c,'分别 ping 目标 IP、主域名与别名。随后进入 Simulation 模式，读取实际 DNS 响应 PDU；响应资源记录通过 Packet Tracer 内置脚本接口提取，截图保留其真实字段。')

_,c=page('数据记录\n和计算','二 初始状态的 IP 与域名测试')
pic(c,'pt-initial-ip-main.jpg','任务2 初始 IP 与主域名 ping 输出',(20,150,1170,840))
qa(c,1,'初始三个 ping 结果如何 DNS 响应中包含哪些资源记录','IP、www.yizhixing.com 与 www1.2023141220023.com 均能到达192.168.1.200。正式三项测试均为4发4收、0%丢包，见图2、图3。刚连线后的预热 IP ping 曾出现1次超时，随后3次回应；链路稳定后重新测试为4发4收，预热记录单独保留。DNS 主域名响应有1条A记录；别名响应有1条CNAME和1条A记录，见图4、图5。')
_,c=page('（接上）\n数据记录','二 初始状态的别名测试')
pic(c,'pt-initial-alias.jpg','任务2 别名解析并 ping 到目标主机',(20,150,1170,530))
para(c,'CNAME 记录把别名指向规范名称，最终 A 记录给出主机 IPv4 地址。别名和规范名称的拼写必须与配置一致，CNAME 的值应填写规范域名。')
para(c,'初始验证汇总：',True)
para(c,'ping 192.168.1.200：4发4收。\nping www.yizhixing.com：4发4收。\nping www1.2023141220023.com：4发4收。')

_,c=page('（接上）\n数据记录','三 Simulation 主域名 DNS 响应')
pic(c,'pt-dns-main-pdu.jpg','任务2 主域名实际 DNS 响应 PDU',(20,173,790,1145),width=3.65)
para(c,'响应事务 ID 为51759，问题名www.yizhixing.com；Answers=1，A=192.168.1.200，Class=IN，TTL=86400秒。此处的 FrameInstance 是仿真内部索引，不能与真实抓包的 Wireshark 帧号混用。')
_,c=page('（接上）\n数据记录','三 Simulation 别名 DNS 响应')
pic(c,'pt-dns-alias-pdu.jpg','任务2 别名实际响应中的 CNAME 与 A 记录',(20,525,850,1168),width=4.2)
para(c,'响应事务 ID 为63948，问题名www1.2023141220023.com；Answers=2。第一条为CNAME，规范名称www.yizhixing.com；第二条为该规范名称的A记录192.168.1.200。两条记录均为Class=IN、TTL=86400秒。')

_,c=page('（接上）\n实验内容','四 DNS 服务器改址故障')
pic(c,'pt-fault.jpg','任务2 DNS 改为 .198 而客户端仍指向 .199',(20,150,1170,665))
qa(c,2,'DNS 改址后结果有何变化 原因是什么 如何恢复','把DNS服务器IP从192.168.1.199改为192.168.1.198，PC仍使用旧解析器.199。直接ping .200仍为4发4收；ping主域名和别名均提示“Ping request could not find host”。原因是名称查询无法到达当前DNS服务地址，而IP通信链路仍可用。将PC的DNS改为.198后，三项测试全部恢复为4发4收，见图7。此次未额外清除仿真DNS缓存；故障是实际重复测试观察到的结果。')
_,c=page('（接上）\n数据记录','四 修复后复测结果')
pic(c,'pt-recovery.jpg','任务2 修复客户端 DNS 后三项测试全部成功',(20,150,1150,1145))
para(c,'工程最终保存为恢复后的正确配置：DNS服务器192.168.1.198/24，PC解析器192.168.1.198；A和CNAME记录保持不变。观察初始响应PDU时曾暂时恢复.199用于重现，再回到.198保存。')

_,c=page('（接上）\n实验内容','五 真实 DNS 捕获与操作顺序')
para(c,'正式抓包：evidence/final/2023141220023_易治行_实验2_DNS.pcapng，共51个分组。下文真实报文帧号均对应这一文件。Wireshark截图中final-dns.pcapng为同一文件的英文名副本，字节与SHA256完全一致。')
para(c,'在en0上用Wireshark自带dumpcap捕获，过滤器为port 53 or icmp，包含查询和响应两个方向。实验操作、起止时间和命令输出保存在real-actions.json。')
para(c,'每个域名先执行以下两条命令清本机DNS缓存，等待2秒；首次ping结束后立即重复同一命令，期间不再清缓存。')
mono(c,'sudo dscacheutil -flushcache\nsudo killall -HUP mDNSResponder\nping -c 3 -i 0.3 -W 1000 www.baidu.com\nping -c 3 -i 0.3 -W 1000 www.scu.edu.cn')
para(c,'两域名分别完成首次与重复ping后，执行dscacheutil -cachedump -entries Host检查缓存；本机返回“Unable to get details from the cache node”，因此缓存结论结合完整抓包与重复命令判断。')
para(c,'四次nslookup均指定202.115.39.9，域名末尾保留点，以完整域名查询。')
for domain in ['www.baidu.com.','www.scu.edu.cn.']:
    for typ in ['A','CNAME']:mono(c,f'nslookup -type={typ} {domain} 202.115.39.9')
para(c,'系统配置显示校园DNS为202.115.39.9和202.115.39.6，但ping实际抓到的端点为223.5.5.5。报告按实际端点分析，不能仅根据系统设置替代报文证据。')

_,c=page('（接上）\n数据记录','六 查询端点与事务匹配')
qa(c,3,'DNS 查询和响应的源目的 IP 传输协议及端口','首次ping百度：第5帧10.135.128.230:54546 →223.5.5.5:53；第6帧方向及端口相反。首次ping川大：第25帧客户端:56898 →223.5.5.5:53，第26帧反向。本次均为IPv4上的UDP。DNS也支持TCP，但本次目标报文没有使用TCP。')
qa(c,4,'事务 ID 是什么 如何将查询与响应对应','DNS头部的16位事务ID用于匹配事务，响应复制请求ID。还应核对反向端点、端口、查询名称/类型与QR位，不能只凭ID判断。百度5↔6为0x12c7，川大25↔26为0xf1c5；Wireshark响应的Request In分别指向5与25。')
para(c,'四组nslookup对应关系如下，客户端IP均为10.135.128.230，服务端为202.115.39.9:53：',True)
for s in ['百度A：36↔37，0x9d7c，客户端端口49533。','百度CNAME：38↔39，0x4552，客户端端口54677。','川大A：40↔41，0x0ce6，客户端端口52277。','川大CNAME：50↔51，0x3a1c，客户端端口53482。']:para(c,s)
para(c,'响应时间示例：百度6−5=0.006491秒，川大26−25=0.020845秒。这里只统计DNS请求到响应的时间，与ping的ICMP往返时延不同。')

_,c=page('（接上）\n数据记录','七 真实响应标志位')
pic(c,'ws-baidu-header.jpg','任务3 第6帧事务 ID 与 AA RD RA 标志',(0,130,1440,1325))
qa(c,7,'客户端发出的 DNS 查询属于递归还是迭代 查询依据是什么','第5、25帧请求RD=1，表示客户端希望解析器递归求解；对应响应RD=1、RA=1，表示服务器声明支持递归，见图8。客户端本次请求属于递归请求。仅在客户端抓包，无法观察解析器与根、顶级及权威服务器的上游交互，不能据此声称已经抓到了完整迭代过程。')

_,c=page('（接上）\n数据记录','八 百度真实响应资源记录')
pic(c,'ws-baidu-response.jpg','任务3 第6帧百度 CNAME 与 A 答案',(0,130,1440,1330))
qa(c,5,'DNS 返回哪些资源记录 是否权威 与 ping 地址有何关系','百度第6帧Answers=3：www.baidu.com CNAME www.a.shifen.com（TTL1029秒）；后者的两条A分别为110.242.70.57与110.242.69.21（TTL均54秒）。AA=0，为非权威响应。ping选择110.242.70.57，属于答案中的地址；首次和重复均3发3收。川大第26帧见图10。')

_,c=page('（接上）\n数据记录','八 四川大学真实响应与别名判断')
pic(c,'ws-scu-response.jpg','任务3 第26帧四川大学 A 响应',(0,130,1440,1330))
para(c,'问题5续：川大第26帧Answers=1，A=202.115.47.167、TTL174秒、AA=0。首次和重复ping均已打印该IP，但各3发0收。DNS解析成功，ICMP未收到回显响应；不能据此判定域名解析失败或服务器离线。')
qa(c,6,'这两个域名是否为别名 规范名称是什么','百度响应中存在www.baidu.com → www.a.shifen.com的CNAME，因此本次百度域名是别名，规范名称为www.a.shifen.com。川大响应直接返回www.scu.edu.cn的A，没有CNAME；任务4的显式CNAME查询也没有返回该类型答案，支持其在本次查询中不是别名。')

_,c=page('（接上）\n数据记录','九 重复查询与缓存证据')
qa(c,8,'首次和重复 ping 时 DNS 报文有何变化 能否判断使用缓存','百度首次ping窗口出现第5/6帧DNS查询/响应，重复窗口只有后续ICMP回显，没有新的目标域名DNS查询；川大首次为第25/26帧，重复窗口也没有新的目标DNS请求。两次重复ping均直接得到与首次相同的目标地址。首次抓到、重复未出现，且间隔未超过本次A记录TTL，支持本机解析结果被复用。')
para(c,'缓存检查命令确实执行，但macOS未给出可枚举的Host缓存明细，因此不能报告一张未看到的缓存表，也不能仅凭本次客户端抓包判断上游DNS服务器使用了哪一级缓存。')
actions=json.loads((E/'evidence/final/real-actions.json').read_text())
for name in ['baidu-ping-first','baidu-ping-repeat','scu-ping-first','scu-ping-repeat']:
    a=next(a for a in actions if a['name']==name)
    stamp=a['start_utc'].split('T')[1][:12]
    para(c,f'{name}：UTC {stamp}；执行{a["end_epoch"]-a["start_epoch"]:.3f}秒。',size=9.5)
para(c,'正式捕获中的其他域名背景DNS流量不计入这两个目标的查询次数。原始捕获保留不删帧；分析使用dns.qry.name等显示过滤条件，并结合操作起止时间。')
para(c,'排障记录：首轮只按校园DNS地址筛选，漏掉ping实际到223.5.5.5的报文；全部接口补抓仅记录到入站响应，无法匹配事务。最终改为en0加双向port 53 or icmp过滤并重新执行，得到完整事务。前两轮单独保留，不用于本报告帧号。')

_,c=page('（接上）\n数据记录','十 A 与 CNAME 查询的比较')
qa(c,9,'A 查询和 CNAME 查询结果有什么不同','A查询请求IPv4地址，查询类型为1；遇到别名时，响应可同时给出CNAME链和最终A地址。CNAME查询请求别名映射，查询类型为5，并不保证返回IPv4地址。百度A响应第37帧有1条CNAME和2条A，CNAME响应第39帧只有1条CNAME。川大A响应第41帧有1条目标A，CNAME响应第51帧答案数为0。')
para(c,'百度第37帧：CNAME TTL615秒，A地址110.242.69.21、110.242.70.57，TTL均233秒。第39帧同一CNAME的TTL为614秒。川大第41帧目标A为202.115.47.167，TTL26秒。解析器和查询时间不同，返回TTL可不同。')
para(c,'第41帧还带有Authority NS和Additional中的名称服务器A地址。这些附加地址不是www.scu.edu.cn的目标A答案，应按资源记录所属区域及Owner Name区分。')
pic(c,'ws-baidu-cname.jpg','任务4 第39帧百度显式 CNAME 查询结果',(0,130,1440,1330))

_,c=page('（接上）\n数据记录','十 显式 CNAME 查询核验')
pic(c,'ws-scu-cname-header.jpg','任务4 第51帧 NOERROR 零答案及 SOA 权威区域',(0,130,1440,1330))
qa(c,10,'两域名的 CNAME 查询分别得到什么 是否支持任务3判断','百度第38/39帧返回www.baidu.com CNAME www.a.shifen.com，支持百度为别名。川大第50/51帧RCODE=0（NOERROR）、Answer RRs=0，Authority中有scu.edu.cn的SOA，主名称服务器dart.scu.edu.cn，SOA TTL300秒；没有CNAME答案。它表示本次无该类型数据，不能写成NXDOMAIN，支持任务3中川大不是别名的判断。')

tt,c=page('结 论\n（结 果）','实验结果')
para(c,'1. Packet Tracer 的A和CNAME记录配置生效，主域名和别名都解析到目标主机192.168.1.200。初始与修复后三项正式ping均为4发4收。')
para(c,'2. DNS服务器改址而客户端不更新时，直接IP通信仍可用，域名与别名解析失败。客户端DNS更新至.198后恢复，说明应分别排查链路连通性与名称解析配置。')
para(c,'3. 真实抓包完成六组目标DNS事务，查询响应ID、端点和名称相互匹配。百度存在CNAME链，川大本次直接返回A；显式CNAME查询复核了这一差异。')
para(c,'4. RD与RA支持客户端请求递归解析的判断。重复ping没有新的目标DNS查询，结合相同地址与TTL支持本机缓存复用；缓存枚举失败和客户端观测范围已在分析中说明。')
def appendrow(index,label):
    el=deepcopy(template.tr_lst[index]);tt._tbl.append(el);row=tt.rows[-1];clean(row)
    for c in [row.cells[0],row.cells[1]]:clear(c)
    para(row.cells[0],label);row.cells[0].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in row.cells[0].paragraphs:fmt(p,align=WD_ALIGN_PARAGRAPH.CENTER)
    return row.cells[1]
c=appendrow(8,'小 结')
para(c,'实验中，先用目标IP确认链路，再检查客户端DNS地址和服务器记录，能够明确定位故障。对抓包应先核对实际接口与端点，再结合事务ID和问题名分析；只有响应而没有查询的捕获不足以完成事务匹配。')
para(c,'分析资源记录时，需要区分Answer、Authority、Additional，也要区分无CNAME数据与域名不存在。ping是否收到ICMP回显不能替代DNS解析结果。')
c=appendrow(9,'指导老师\n评 议');para(c,'\n\n成绩评定：                 指导教师签名：')

doc.core_properties.author='易治行';doc.core_properties.title='实验2 DNS服务器配置与DNS协议分析 实验报告'
out=E/'2023141220023_易治行_实验2_报告.docx';doc.save(out)
assert hashlib.sha256((R/'experiments/实验报告模板.doc').read_bytes()).hexdigest()==hsh
print(json.dumps({'docx':str(out),'figures':fig},ensure_ascii=False))
