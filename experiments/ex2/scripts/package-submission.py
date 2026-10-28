from pathlib import Path
import hashlib, json, shutil, zipfile
import pdfplumber
from docx import Document

E=Path(__file__).resolve().parents[1]
R=E.parents[1]
stem='2023141220023_易治行_实验2'
render=R/'tmp/pdfs/ex2/report/render'
pdf=E/'output/pdf'/f'{stem}_报告.pdf'
pdf.parent.mkdir(parents=True,exist_ok=True)
shutil.copy2(render/pdf.name,pdf)
docx=E/f'{stem}_报告.docx'
doc=Document(docx)
with zipfile.ZipFile(docx) as z:assert z.testzip() is None
assert len(doc.inline_shapes)==12
alltext='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(p.text for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs)
for s in ['2023141220023','易治行','102班','192.168.1.198','0x12c7','0xf1c5','NOERROR','NXDOMAIN']:
    assert s in alltext,s
assert all(f'{i}. ' in alltext for i in range(1,11))
assert all(f'图{i} ' in alltext for i in range(1,13))
for s in ['TODO','待补充','尚未完成','待填写']:assert s not in alltext,s
with pdfplumber.open(pdf) as p:
    assert len(p.pages)==17
    assert all(len(pg.extract_text() or '')>80 for pg in p.pages)
cap=E/'evidence/final'/f'{stem}_DNS.pcapng'
assert cap.stat().st_size==7164
assert hashlib.sha256(cap.read_bytes()).digest()==hashlib.sha256((R/'tmp/pdfs/ex2/final-dns.pcapng').read_bytes()).digest()
records=json.loads((E/'evidence/final-analysis/analysis.json').read_text())
by={x['frame']:x for x in records}
for q,a in [(5,6),(25,26),(36,37),(38,39),(40,41),(50,51)]:
    assert by[q]['id']==by[a]['id']
    assert by[q]['src']==by[a]['dst'] and by[q]['dst']==by[a]['src']
    assert by[a]['response_to']==str(q)
assert by[51]['rcode']=='0' and by[51]['answers']=='0'
assert (E/f'{stem}.pkt').stat().st_size>60000
verification={'pdf_pages':17,'figures':12,'analysis_questions':10,'docx_zip_integrity':'passed','visual_qa':'all 17 retained pages verified after removing final notes page','formal_capture_packets':51,'formal_capture_bytes':cap.stat().st_size,'target_dns_transaction_pairs':6,'capture_english_copy_sha256_matches':True,'packet_tracer_file_bytes':(E/f'{stem}.pkt').stat().st_size}
(E/'evidence/verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2))
files=[E/'README.md',docx,pdf,E/f'{stem}.pkt',E/'实验2_DNS服务器配置与DNS协议分析_实验指导书.pdf',E/'evidence/verification.json']
files+=sorted((E/'evidence/final').glob('*'))
files+=sorted((E/'evidence/final-analysis').glob('*'))
files+=sorted((E/'evidence').glob('pt-*.jpg'))+sorted((E/'evidence').glob('ws-*.jpg'))
files+=sorted(p for p in (E/'scripts').iterdir() if p.suffix in ['.py','.sh','.js'])
files=[p for p in files if p.is_file()]
manifest={str(p.relative_to(E)):hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
mp=E/'evidence/submission-sha256.json';mp.write_text(json.dumps(manifest,ensure_ascii=False,indent=2));files.append(mp)
out=E/f'{stem}_提交材料.zip'
with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
    for p in files:z.write(p,str(p.relative_to(E)))
with zipfile.ZipFile(out) as z:
    assert z.testzip() is None
    for n,d in manifest.items():assert hashlib.sha256(z.read(n)).hexdigest()==d,n
print(json.dumps({'pdf':str(pdf),'zip':str(out),'members':len(files),'zip_bytes':out.stat().st_size,'checks':verification},ensure_ascii=False,indent=2))
