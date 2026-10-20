from pathlib import Path
import json
from docx import Document
from PIL import Image
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
R=Path(__file__).resolve().parents[1]
D=json.loads((R/'evidence/report-data.json').read_text())
f=D['frames']; base='http://gaia.cs.umass.edu/wireshark-labs/'
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 s=doc.styles[name]; s.font.name='Times New Roman'; s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
 s.font.size=Pt(10.5)
 s.paragraph_format.space_after=Pt(5); s.paragraph_format.line_spacing=1.05
for name,size in [('Title',19),('Heading 1',14),('Heading 2',12)]:
 doc.styles[name].font.size=Pt(size); doc.styles[name].font.bold=True
 doc.styles[name].paragraph_format.keep_with_next=True
for st in doc.styles:
 if st.type == 1 or st.type == 2:
  st.font.name='Arial Unicode MS'
  rf=st.element.get_or_add_rPr().get_or_add_rFonts()
  for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Arial Unicode MS')
  for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
   rf.attrib.pop(qn('w:'+attr),None)
  pp=st.element.find(qn('w:pPr'))
  if pp is not None:
   for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.author='易治行'; doc.core_properties.title='实验1 Web服务器搭建与HTTP协议分析'

def p(t,style=None): return doc.add_paragraph(t,style)
pending_page=False
def h(t):
 global pending_page
 x=doc.add_heading(t,1)
 if pending_page:x.paragraph_format.page_break_before=True;pending_page=False
def qa(n,q,a):
 x=p(f'{n}．{q}'); x.paragraph_format.keep_with_next=True
 for r in x.runs:r.bold=True
 p(a)
def pic(file,cap,width=6.6):
 x=doc.add_paragraph(); x.paragraph_format.keep_with_next=True
 source=R/'evidence'/file
 assetdir=R.parents[1]/'tmp/pdfs/ex1/report-assets';assetdir.mkdir(parents=True,exist_ok=True)
 target=assetdir/file
 with Image.open(source) as im:
  clean=Image.frombytes(im.mode,im.size,im.tobytes());clean.save(target,format='PNG')
  with Image.open(target) as restored:assert restored.tobytes()==im.tobytes()
 x.add_run().add_picture(str(target),width=Inches(width))
 c=p(cap,'Caption'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].font.size=Pt(9)
def page():
 global pending_page
 pending_page=True
def table(rows):
 t=doc.add_table(rows=0,cols=2);t.style='Table Grid';t.autofit=False
 t.columns[0].width=Inches(1.35); t.columns[1].width=Inches(5.42)
 for a,b in rows:
  c=t.add_row().cells;c[0].text=a;c[1].text=b
  c[0].width=Inches(1.35);c[1].width=Inches(5.42)
  for cell in c:
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(3)
 return t
p('实验1 Web服务器搭建与HTTP协议分析','Title')
p('姓名：易治行    学号：2023141220023    班级：102班\n实验日期：2026年10月19日')
p('本实验在 macOS 上启用 Apache 静态网站，使用浏览器访问指定 HTTP 资源，并通过 Wireshark 工具链捕获、分析请求与响应。实验分别观察本地缓存复用与刷新时的条件请求，确认是否重新传输了页面正文。')
h('一 实验环境与服务器配置')
table([('操作系统','macOS 26.5.1，Apple Silicon'),('Web服务器','Apache HTTP Server 2.4.66，系统自带 /usr/sbin/httpd'),('浏览器','Chromium 153.0.0.0 独立临时会话；Safari 26.5 另作首页访问验证'),('抓包与分析','Wireshark dumpcap 4.6.9 正式捕获；TShark 4.6.9 解析；Wireshark 4.4.14 显示截图'),('获取与启用','Apache 随 macOS 提供；Wireshark 从官方下载目录获取。Apache 使用实验目录内独立配置启动。'),('监听与访问','127.0.0.1:8080；http://127.0.0.1:8080/'),('默认首页','index.html'),('网站根目录',str(R/'site')),('服务器配置文件',str(R/'apache/httpd.conf'))])
p('启动前运行 httpd -t -f 配置文件，输出 Syntax OK；随后启动 httpd，并用 curl 和浏览器实际访问根路径。响应为 HTTP/1.1 200 OK，页面显示姓名、学号、班级与实验名称。')
p('URL 中的 127.0.0.1 指向本机，8080 为 TCP 监听端口；请求路径 / 对应 DocumentRoot 所指目录，再由 DirectoryIndex 选择 index.html。若访问 /index.html，则直接映射到同一个首页文件。使用非默认 HTTP 端口时，URL 必须显式带上 :8080。')
page();h('二 网站运行结果')
pic('01-apache-browser.png','图1 浏览器通过 HTTP 访问 Apache 默认首页',6.5)
p('DocumentRoot 决定 URL 路径对应的本地目录，DirectoryIndex 决定目录请求的默认文档。Require all granted 允许读取该网站目录；监听地址仅为本机回环地址，因此本次验证的是本机访问。浏览器与服务器可以运行在同一台计算机上。')
page();h('三 服务器配置与抓包方法')
pic('08-apache-config.png','图2 实际加载的 Apache 配置文件',6.55)
p('Listen、DocumentRoot 和 DirectoryIndex 分别对应监听端点、网站目录与默认首页。完整配置保存于 apache/httpd.conf，语法检查结果为 Syntax OK。')
p('正式捕获使用 Wireshark 自带 dumpcap，在 en0 接口应用捕获过滤器 host gaia.cs.umass.edu and tcp port 80。该捕获仅覆盖指定实验站点的明文 HTTP。浏览器采用新的临时会话并使用 --no-proxy-server，仅该实验会话直连，不修改系统代理。')
p('原始文件为 evidence/wireshark-browser.pcapng。图3至图8均来自该正式捕获；共39个分组，帧号与原文件对应。先前为验证权限使用 tcpdump 生成的文件单独保留为预实验，不作为下文正式答案的数据来源。')
page();h('四 基本 HTTP 请求与响应分析')
p('测试地址：'+base+'HTTP-wireshark-file1.html')
pic('02-http-request.png',f'图3 基本 GET 请求 第{f["basic_req"]}帧',6.55)
qa(1,'请求方法、资源与 HTTP 版本是什么？',f'图3第{f["basic_req"]}帧请求行为 GET /wireshark-labs/HTTP-wireshark-file1.html HTTP/1.1，方法是 GET，资源为上述 HTML 路径，版本是 HTTP/1.1。Host 为 gaia.cs.umass.edu。')
qa(2,'响应版本、状态码与状态短语是什么？',f'图4第{f["basic_res"]}帧状态行为 HTTP/1.1 200 OK，版本为 HTTP/1.1，状态码为200，状态短语为 OK。')
qa(3,'客户端、服务器的 IP 与 TCP 端口分别是什么？',f'图3的 IPv4 与 TCP 摘要显示：客户端 {D["client"]}:{D["basic_port"]}，服务器 {D["server"]}:80；图4响应方向相反，证明两者属于同一连接。客户端地址是抓包接口上看到的局域网地址，不将其当成公网出口地址。')
page();h('四 基本 HTTP 请求与响应分析 续')
qa(4,'Accept-Language 包含什么语言偏好？','图3显示 Accept-Language: zh-CN，表示客户端偏好中国大陆使用的中文。该值来自本次独立浏览器会话的语言设置，不代表服务器页面实际使用的语言。')
qa(5,'Accept-Encoding 与实际编码有什么区别？','图3显示 Accept-Encoding: gzip, deflate，表示浏览器能够接受的内容编码。图4未出现 Content-Encoding，且正文可直接读到 HTML，说明本次响应没有通过该首部声明 gzip 或 deflate 等内容编码；不能从请求字段推断响应实际被压缩。UTF-8 是字符编码，也不等同于 gzip 内容编码。')
qa(6,'Accept 与 Content-Type 的含义及区别是什么？','图3的 Accept 包含 text/html、application/xhtml+xml、application/xml;q=0.9，以及 image/avif、image/webp、image/apng、*/*;q=0.8 和 application/signed-exchange;v=b3;q=0.7。它表达客户端可接受的媒体类型及相对偏好；图4的 Content-Type: text/html; charset=UTF-8 则说明服务器实际返回的是 UTF-8 编码的 HTML。')
pic('03-http-response.png',f'图4 基本响应 第{f["basic_res"]}帧',6.55)
page();h('四 基本 HTTP 请求与响应分析 续二')
qa(7,'资源最后修改时间通过什么字段获得？','图4显示 Last-Modified: Tue, 28 Oct 2025 05:59:01 GMT，因此服务器声明的最后修改时间为2025年10月28日05:59:01 UTC，即北京时间13:59:01。Date 表示本次响应生成的时间，不能替代 Last-Modified。')
qa(8,'响应内容长度是多少，是报文总长还是实体长度？',f'图4的 Content-Length: 128 表示响应内容为128字节，不包含 HTTP 状态行与首部，也不是 Ethernet/IP/TCP 帧的总长度。第{f["basic_res"]}帧的 TCP 负载长度为{D["basic_tcp_len"]}字节，包含 HTTP 首部与正文，因而大于128字节。')
qa(9,'如何确认响应属于目标网页？',f'图3请求 URI 是 /wireshark-labs/HTTP-wireshark-file1.html，Host 是 gaia.cs.umass.edu；图4的 Request in frame 指向第{f["basic_req"]}帧，正文中也含有目标 file1.html 的下载成功文字。与之区分，第{f["favicon_req"]}帧请求 /favicon.ico，第{f["favicon_res"]}帧响应301，两者不是目标 HTML 的请求响应。')
qa(10,'连接是否允许保持，仅一组报文能否证明复用？',f'请求和响应均为 HTTP/1.1，且包含 Connection: keep-alive，响应还有 Keep-Alive: timeout=5, max=100，说明允许连接保持。一组请求响应只能证明这次交换，不能单独证明复用；本次进一步检查同一 tcp.stream={D["basic_stream"]}，发现第{f["basic_req"]}帧的页面 GET 与第{f["favicon_req"]}帧的 favicon GET 使用同一客户端端口 {D["basic_port"]}，中间没有建立新的 TCP 连接，因此整段抓包确实提供了连接复用证据。')
h('五 Conditional GET 操作过程')
p('测试地址：'+base+'HTTP-wireshark-file2.html')
table([('首次访问',D['first_time']+f'；第{f["first_req"]}帧 GET，第{f["first_res"]}帧200；正文371字节'),('再次打开同一 URL',D['second_time']+'；以同一地址再次导航，未出现新的目标 HTTP 请求；浏览器导航计时 transferSize=0，encodedBodySize=371'),('补充普通刷新',D['reload_time']+f'；第{f["cond_req"]}帧条件GET，第{f["cond_res"]}帧304；无正文')])
p('首次访问前使用全新临时浏览器会话，目标资源无既有缓存。首次加载后立即重复访问同一 URL，随后额外进行一次普通刷新；未使用硬刷新，也未手动构造条件首部。下文将第二次普通访问与补充刷新分开解释。')
page();h('六 Conditional GET 分析问题')
qa(1,'第一次 GET 是否包含缓存验证字段？',f'图5第{f["first_req"]}帧中没有 If-Modified-Since 或 If-None-Match。该会话此前没有目标资源缓存，因而没有已保存的验证器可以附带，首次采用无条件 GET；图6第{f["first_res"]}帧返回200和371字节正文。')
qa(2,'第一次响应包含哪些缓存验证字段？','图6包含 Last-Modified: Tue, 28 Oct 2025 05:59:01 GMT，以及 ETag: "173-64231b67176b7"。前者提供资源最后修改时间，可用于后续 If-Modified-Since；后者提供该表示的实体标签，可用于 If-None-Match。两者都来自实际响应。')
qa(3,'第二次访问是否发出了 GET，验证器如何关联？',f'第二次通过同一 URL 再次导航时，没有捕获到新的目标 GET；同时浏览器导航计时显示 transferSize=0、encodedBodySize=371，支持直接使用本地缓存。为排除抓包接口失效、HTTPS升级、代理路径变化与请求被漏选，核查了同一捕获内的首次明文交换及随后刷新的明文交换、目标80端口过滤器、直连启动参数和实际导航URL。补充普通刷新时，第{f["cond_req"]}帧才出现 GET，If-None-Match 与首次响应的 ETag 完全相同，If-Modified-Since 与首次响应的 Last-Modified 完全相同，见图7。')
qa(4,'为何未出现请求或返回304，对传输有什么影响？',f'第二次普通访问没有发生本次网络缓存验证，也未重新传输正文；首次响应未给出显式 Cache-Control/Expires，但具有较早的 Last-Modified，允许浏览器采用启发式新鲜度，本地缓存命中与此相符，具体内部新鲜度算法未由抓包证明。补充刷新请求包含 Cache-Control: max-age=0 及两个条件字段，服务器在第{f["cond_res"]}帧返回304 Not Modified，确认现有表示可继续使用，未附带正文。由此避免再次传输371字节正文；304仍有首部和TCP等开销，不能把节省的正文长度等同于全部网络字节节省量。')
p('当 If-None-Match 与 If-Modified-Since 同时出现时，RFC 9110规定优先评估 If-None-Match。本次只能观察验证器与304的对应关系，无法直接看到服务器内部比较代码。')
h('七 资料查阅与排障')
p('权限问题：抓包设备 /dev/bpf0 的所有者为 root，当前账户直接执行 tcpdump 得到 Permission denied。先以管理员权限运行定向 tcpdump 验证；Wireshark 安装完成后，再由本人在终端启动其自带 dumpcap 正式捕获。成功生成 pcapng，并由 Wireshark/TShark 识别请求、响应及条件字段，证明捕获链路有效。未修改全局 BPF 权限。')
p('显示问题：Wireshark 4.6.9 的界面在自动读取辅助功能结构时出现 EXC_BAD_ACCESS，TShark 可正常解析同一文件。改用官方 Wireshark 4.4.14 查看原始 pcapng，界面稳定；截图的帧号和载荷与4.6.9解析结果一致。')
page();h('八 首次访问的请求与响应证据')
pic('04-cache-first-request.png',f'图5 首次 GET 第{f["first_req"]}帧 没有缓存验证字段',6.55)
p('图5显示完整的请求首部区与结束空行。首次请求没有 If-None-Match 和 If-Modified-Since；因此后续条件字段不是首次请求就具备的。')
p('首次响应的状态、正文长度与两个验证器见图6。本次首次与刷新的客户端端口同为62842，并复用了原连接。缓存复用与 TCP 连接复用属于不同机制，前者决定是否使用已保存表示，后者决定是否沿用已建立连接。')
page();h('八 首次访问的请求与响应证据 续')
pic('05-cache-first-response.png',f'图6 首次200响应 第{f["first_res"]}帧 371字节正文与两个验证器',6.55)
p('图6中的实体标签为 "173-64231b67176b7"，最后修改时间为 Tue, 28 Oct 2025 05:59:01 GMT。浏览器保存正文后，后续刷新能够把这些值作为条件发回服务器。')
p('再次普通导航的浏览器计时证据为 transferSize=0、encodedBodySize=371、decodedBodySize=371；对应完整记录保存于 evidence/final-second-visit-cache.txt。')
page();h('九 普通刷新的条件 GET 证据')
pic('06-conditional-request.png',f'图7 补充刷新请求 第{f["cond_req"]}帧',6.55)
p('图7同时出现 Cache-Control: max-age=0、If-None-Match: "173-64231b67176b7" 和 If-Modified-Since: Tue, 28 Oct 2025 05:59:01 GMT。两个验证器分别与图6的响应字段一致。')
p('这次操作是首次加载和第二次缓存导航之后的普通刷新，不把它写成第二次普通访问的网络响应。')
page();h('九 普通刷新的条件 GET 证据 续')
pic('07-conditional-response.png',f'图8 304响应 第{f["cond_res"]}帧 没有正文',6.55)
p('图8状态行为 HTTP/1.1 304 Not Modified，Request in frame 指向图7的请求。HTTP首部结束后没有 File Data 或 HTML 正文，TCP负载为首部本身；本次没有出现 Content-Length，不补写 Content-Length: 0。')
page();h('十 结果与简短小结')
qa(1,'依据哪些配置或抓包证据确认或修正原有判断？','Apache 配置中的 DocumentRoot、DirectoryIndex 与浏览器根路径访问结果共同确认了默认首页映射。再次普通访问没有网络请求，而补充刷新才出现条件GET和304，这修正了“重复打开页面必然返回304”的判断。')
qa(2,'若供他人使用，重点核验什么，资料查阅在哪一步有帮助？','应核验监听地址是否包含可达网卡、访问控制和防火墙是否允许预期用户访问，并从另一台计算机验证实际URL与页面身份。本次查阅Apache官方的DocumentRoot与DirectoryIndex文档帮助完成路径映射；HTTP缓存规范帮助区分本地新鲜缓存和304验证，查明重复访问未发请求的解释边界。')
page();h('十一 参考资料与材料索引')
for t in [
'Apache Core Features，DocumentRoot：https://httpd.apache.org/docs/2.4/mod/core.html#documentroot 。用于解释网站根目录与URL映射。',
'Apache mod_dir，DirectoryIndex：https://httpd.apache.org/docs/2.4/mod/mod_dir.html#directoryindex 。用于配置根路径默认首页。',
'RFC 9110 HTTP Semantics，第8.8节、第13.1节、第15.4.5节：https://httpwg.org/specs/rfc9110.html 。用于解释验证器、条件首部和304。',
'RFC 9111 HTTP Caching，第4.2.2节：https://httpwg.org/specs/rfc9111.html#heuristic.freshness 。用于解释未显式指定过期时间时的启发式缓存。',
'Wireshark User’s Guide：https://www.wireshark.org/docs/wsug_html_chunked/ 。用于捕获文件读取、显示过滤器和报文详情分析。',
'Wireshark官方下载目录：https://www.wireshark.org/download/osx/all-versions/ 。本次获取4.6.9与4.4.14 Arm 64安装镜像。']:
 p(t)
p('查阅日期均为2026年10月19日。协议资料用于解释本次现象；IP、端口、状态码和具体首部值均取自本次正式抓包。')
table([('正式原始抓包','evidence/wireshark-browser.pcapng'),('可读报文全文','evidence/final-http-details.txt'),('结构化解析','evidence/final-http.json'),('浏览器导航记录','evidence/final-actions.json 与 final-second-visit-cache.txt'),('截图','evidence/01-apache-browser.png 至 08-apache-config.png'),('服务器配置','apache/httpd.conf'),('自建网页','site/index.html'),('启停脚本','scripts/server.sh'),('重现实验说明','README.md')])
p('材料中的文件路径以本实验目录为基准。报告截图保留原始抓包帧号；复核时可在Wireshark显示过滤器中输入 frame.number == 帧号。')
if not D.get('formal',False):
 replacements={
 '通过 Wireshark 工具链捕获、分析请求与响应':'捕获请求与响应，并通过 Wireshark 分析',
 'Wireshark dumpcap 4.6.9 正式捕获；TShark 4.6.9 解析；Wireshark 4.4.14 显示截图':'macOS tcpdump 捕获；TShark 4.6.9 解析；Wireshark 4.4.14 显示截图',
 '正式捕获使用 Wireshark 自带 dumpcap，在 en0 接口应用捕获过滤器':'本次预实验使用 macOS 自带 tcpdump，以 pktap,all 捕获并在 en0 观察到流量；捕获过滤器为',
 '原始文件为 evidence/wireshark-browser.pcapng。图3至图8均来自该正式捕获；共39个分组，帧号与原文件对应。先前为验证权限使用 tcpdump 生成的文件单独保留为预实验，不作为下文正式答案的数据来源。':'原始文件为 evidence/browser-http.pcap，实际格式为pcapng；分析快照为 evidence/http-analysis.pcapng。图3至图8均为本机预实验的实际报文。由于系统权限限制，尚未完成由Wireshark自带dumpcap执行的正式捕获，因此该步骤待补；不能将此预实验视为严格完成全部指导书要求。',
 '先以管理员权限运行定向 tcpdump 验证；Wireshark 安装完成后，再由本人在终端启动其自带 dumpcap 正式捕获。成功生成 pcapng':'以管理员权限运行定向 tcpdump 验证，成功生成 pcapng；Wireshark自带dumpcap的正式捕获仍待本人在终端启动。预实验文件',
 '图3至图8均来自该正式捕获':'图3至图8均来自该预实验捕获',
 '查看原始 pcapng':'查看分析快照 pcap',
 '本次正式抓包':'本次实际预实验抓包',
 '正式原始抓包':'预实验原始抓包',
 'evidence/wireshark-browser.pcapng':'evidence/browser-http.pcap',
 'evidence/final-http-details.txt':'evidence/http-details.txt',
 'evidence/final-http.json':'evidence/http.json',
 'evidence/final-actions.json 与 final-second-visit-cache.txt':'evidence/second-visit-cache.txt',
 }
 allp=list(doc.paragraphs)+[pp for t in doc.tables for row in t.rows for c in row.cells for pp in c.paragraphs]
 for pp in allp:
  new=pp.text
  for a,b in replacements.items():new=new.replace(a,b)
  if new!=pp.text:pp.text=new
 name='实验1报告_待补Wireshark正式捕获_易治行.docx'
else:
 name='实验1报告_易治行_2023141220023.docx'
out=R/name;doc.save(out);print(out)
